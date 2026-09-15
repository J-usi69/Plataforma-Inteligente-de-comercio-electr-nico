# -*- coding: utf-8 -*-
"""Genera el .docx con los Casos de Prueba 17-34 (CUs que faltaban en el documento del parcial)."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

CASES = [
    (17, "Gestión de Roles y Permisos", "CU-04: Gestionar roles y permisos",
     "Verificar que el administrador pueda crear roles, configurar sus permisos por módulo y asignarlos a usuarios.",
     "El administrador se encuentra autenticado con permisos de configuración del sistema.",
     [
        ("Acceder a \u201cRoles y Permisos\u201d.", "El sistema muestra los roles existentes con sus permisos por módulo."),
        ("Crear un nuevo rol e indicar los permisos por módulo.", "El sistema acepta los datos ingresados."),
        ("Guardar el rol.", "El rol y sus permisos quedan almacenados correctamente."),
        ("Asignar el nuevo rol a un usuario.", "El usuario queda vinculado al rol asignado."),
        ("Intentar eliminar un rol con usuarios activos asignados.", "El sistema impide la eliminación y solicita reasignar previamente a esos usuarios."),
     ], "Pendiente"),

    (18, "Gestión de Personal", "CU-05: Gestionar personal",
     "Validar que el administrador pueda registrar, editar y desactivar personal (Encargado/Cajero) de una sucursal.",
     "El administrador está autenticado y las sucursales ya se encuentran registradas (CU-06).",
     [
        ("Acceder a \u201cGestionar personal\u201d.", "El sistema muestra el listado de personal existente."),
        ("Registrar personal indicando nombres, cargo (Encargado o Cajero) y sucursal asignada.", "El sistema valida los datos ingresados."),
        ("Guardar el registro.", "Se crea el Usuario base (si no existe) y el registro Personal vinculado, en estado activo."),
        ("Editar los datos de un miembro de personal existente.", "Los cambios quedan almacenados correctamente."),
        ("Registrar personal con un correo/celular ya existente como Usuario.", "El sistema solicita vincular la cuenta existente en lugar de duplicarla."),
     ], "Pendiente"),

    (19, "Gestión de Proveedores", "CU-07: Gestionar proveedores",
     "Comprobar el registro, edición y desactivación de proveedores de la cadena.",
     "El administrador se encuentra autenticado con permisos de gestión de proveedores.",
     [
        ("Acceder a \u201cGestionar proveedores\u201d.", "El sistema muestra el listado de proveedores existentes."),
        ("Registrar un nuevo proveedor indicando nombre de empresa y contacto.", "El sistema valida los datos ingresados."),
        ("Guardar el proveedor.", "El proveedor queda registrado en estado activo."),
        ("Editar un proveedor existente.", "Los cambios quedan almacenados correctamente."),
        ("Desactivar un proveedor con prendas o movimientos de inventario activos asociados.", "El sistema advierte del impacto antes de confirmar la baja."),
     ], "Pendiente"),

    (20, "Gestión de Temporadas y Colecciones", "CU-10: Gestionar temporadas y colecciones",
     "Validar el registro de temporadas y colecciones, incluyendo la coherencia de sus fechas.",
     "El administrador está autenticado con permisos de catálogo.",
     [
        ("Acceder a \u201cTemporadas y colecciones\u201d.", "El sistema muestra el listado existente."),
        ("Registrar una nueva temporada con fechaInicio anterior a fechaFin.", "El sistema valida la coherencia de las fechas."),
        ("Guardar la temporada.", "La temporada queda registrada en estado activo."),
        ("Registrar una colección asociada a la temporada creada.", "La colección queda vinculada correctamente a la temporada."),
        ("Registrar una temporada con fechaInicio posterior a fechaFin.", "El sistema rechaza o advierte la inconsistencia antes de guardar."),
     ], "Pendiente"),

    (21, "Reserva de Prendas", "CU-15: Reservar prendas",
     "Verificar que el cliente pueda reservar una o varias prendas en una sucursal para probárselas físicamente.",
     "El cliente ha iniciado sesión y existe stock disponible de las variantes elegidas en la sucursal seleccionada.",
     [
        ("Agregar una o varias variantes de prenda a la selección de reserva.", "Las variantes quedan agregadas a la selección."),
        ("Elegir la sucursal y un horario de atención aproximado.", "El sistema acepta la selección."),
        ("Confirmar la reserva.", "El sistema valida que exista stockDisponible suficiente para cada variante en esa sucursal."),
        ("Crear la reserva.", "Se registra la Reserva en estado \u201cpendiente\u201d con sus detalles y se incrementa el stockReservado."),
        ("Intentar reservar una variante sin stock disponible.", "El sistema informa el conflicto y solicita ajustar la selección antes de crear la reserva."),
     ], "Satisfactorio"),

    (22, "Consulta y Cancelación de Reservas", "CU-16: Consultar / cancelar reserva",
     "Comprobar que el cliente pueda revisar el estado de sus reservas y cancelar una reserva pendiente.",
     "El cliente cuenta con al menos una Reserva registrada.",
     [
        ("Acceder a \u201cMis reservas\u201d.", "El sistema muestra las reservas del cliente con su estado y detalle."),
        ("Seleccionar una reserva pendiente.", "Se muestra el detalle de la reserva seleccionada."),
        ("Cancelar la reserva.", "El estado de la Reserva cambia a \u201ccancelada\u201d y se libera el stockReservado."),
        ("Consultar una reserva ya atendida.", "El sistema permite únicamente consultarla, sin opción de cancelar."),
        ("Intentar cancelar una reserva ya atendida.", "El sistema no permite la cancelación."),
     ], "Satisfactorio"),

    (23, "Gestión de Reservas Recibidas en Sucursal", "CU-17: Gestionar reservas recibidas",
     "Validar que el encargado de sucursal pueda visualizar y preparar las reservas asignadas a su sucursal.",
     "Existen reservas pendientes asociadas a la sucursal del encargado. El encargado (o el Cajero, con acceso de solo lectura) se encuentra autenticado.",
     [
        ("Acceder al listado de reservas pendientes de la sucursal.", "El sistema muestra cada Reserva con sus DetalleReserva (variantes y cantidades)."),
        ("Separar físicamente las prendas indicadas en una reserva.", "El encargado confirma la disponibilidad física de las prendas."),
        ("Marcar la reserva como \u201cconfirmada\u201d.", "El sistema actualiza el estado de la Reserva."),
        ("Consultar el detalle de una reserva ya confirmada.", "El detalle refleja el nuevo estado \u201cconfirmada\u201d."),
        ("Reportar una discrepancia (prenda no encontrada físicamente).", "El sistema genera un MovimientoInventario de ajuste."),
     ], "Satisfactorio"),

    (24, "Confirmación de Recepción del Cliente", "CU-18: Confirmar recepción del cliente",
     "Verificar el registro de la llegada del cliente a la sucursal y la entrega de las prendas reservadas.",
     "La reserva del cliente se encuentra en estado \u201cconfirmada\u201d.",
     [
        ("Buscar la reserva del cliente por código o nombre.", "El sistema muestra el detalle de la reserva confirmada."),
        ("Entregar las prendas y marcar la reserva como \u201catendida\u201d.", "El estado de la Reserva cambia correctamente a \u201catendida\u201d."),
        ("Verificar que la reserva quede disponible como reservaOrigen.", "La reserva queda habilitada para originar una venta (CU-19)."),
        ("Marcar como \u201cno show\u201d una reserva fuera del horario de atención previsto.", "El sistema libera el stockReservado correspondiente."),
        ("Intentar confirmar la recepción de una reserva que aún no está \u201cconfirmada\u201d.", "El sistema rechaza la operación."),
     ], "Satisfactorio"),

    (25, "Registro de Venta Presencial", "CU-19: Registrar venta presencial",
     "Validar el registro de una venta presencial a partir de las prendas que el cliente decide comprar en la sucursal.",
     "El cliente ha probado las prendas y decidido cuáles comprar; existe stock suficiente en la sucursal.",
     [
        ("Seleccionar las prendas/variantes que el cliente decidió comprar.", "Las prendas quedan agregadas a la venta."),
        ("Calcular precioUnitario, subtotal y total de la venta.", "El sistema calcula correctamente los montos a partir de VariantePrenda/Prenda."),
        ("Confirmar la venta.", "Se crea la Venta (tipoOrigen=presencial) en estado \u201cpendiente de pago\u201d, vinculada a la reservaOrigen si corresponde."),
        ("Continuar al proceso de pago en caja.", "El sistema habilita el paso al cobro (CU-20)."),
        ("Intentar vender una prenda ya sin stock disponible por venta simultánea en otro canal.", "El sistema notifica el conflicto antes de confirmar."),
     ], "Satisfactorio"),

    (26, "Procesamiento de Pago en Caja", "CU-20: Procesar pago en caja",
     "Comprobar el cobro de una venta presencial en el punto de caja y la actualización automática del inventario.",
     "Existe una Venta registrada en estado \u201cpendiente de pago\u201d.",
     [
        ("Seleccionar el método de pago (efectivo, tarjeta o QR).", "El sistema acepta la selección del método de pago."),
        ("Registrar el monto recibido.", "El sistema valida el monto ingresado."),
        ("Confirmar el pago.", "Se crea el registro Pago con estado=aprobado, vinculado a la Venta, que pasa a estado \u201cpagada\u201d."),
        ("Verificar la actualización automática del inventario tras el pago.", "El sistema dispara la actualización de inventario (CU-25) y genera el comprobante (CU-23)."),
        ("Registrar un pago con monto insuficiente o rechazado.", "La Venta permanece \u201cpendiente de pago\u201d y el sistema permite reintentar."),
     ], "Satisfactorio"),

    (27, "Compra desde Plataforma Web/Móvil", "CU-21: Comprar desde plataforma web/móvil",
     "Verificar que el cliente pueda completar una compra digital de prendas sin acudir presencialmente a una sucursal.",
     "El cliente ha iniciado sesión y existen variantes con stockDisponible en al menos una sucursal.",
     [
        ("Agregar variantes de prenda al carrito de compra.", "El sistema valida el stockDisponible de cada variante en la sucursal de despacho seleccionada."),
        ("Confirmar el carrito.", "El sistema calcula el total de la compra."),
        ("Crear la venta digital.", "Se crea la Venta (tipoOrigen=digital) con sus DetalleVenta, en estado \u201cpendiente de pago\u201d."),
        ("Continuar al proceso de pago electrónico.", "El sistema redirige al cliente al flujo de pago (CU-22)."),
        ("Confirmar el carrito con una variante que quedó sin stock.", "El sistema informa el conflicto y solicita ajustar el carrito antes de continuar."),
     ], "Satisfactorio"),

    (28, "Procesamiento de Pago Electrónico (Stripe)", "CU-22: Procesar pago electrónico (pasarela)",
     "Validar el cobro digital de una venta mediante la pasarela Stripe (modo de prueba) para el método tarjeta, y el flujo simulado para QR.",
     "Existe una Venta digital en estado \u201cpendiente de pago\u201d; la integración de Stripe en modo test está configurada en backend, web y móvil.",
     [
        ("Elegir el método de pago \u201cTarjeta\u201d.", "El sistema crea un PaymentIntent real en Stripe (endpoint crear-intento-pago) y muestra el formulario de tarjeta (Stripe Elements en web / CardField en móvil)."),
        ("Ingresar los datos de la tarjeta de prueba 4242 4242 4242 4242.", "Stripe confirma el pago (PaymentIntent en estado \u201csucceeded\u201d)."),
        ("El sistema verifica el pago contra la API de Stripe y notifica pagar-digital.", "Se registra el Pago (pasarela=Stripe, transaccion_id=id del PaymentIntent) y la Venta pasa a \u201cpagada\u201d."),
        ("Elegir el método de pago \u201cQR\u201d (Libélula).", "El sistema aprueba el pago de forma simulada, sin contactar Stripe, y marca la Venta como \u201cpagada\u201d."),
        ("Intentar notificar un pago con tarjeta sin completar el PaymentIntent en Stripe.", "El sistema rechaza la operación indicando que el pago no se completó."),
     ], "Satisfactorio"),

    (29, "Emisión de Comprobante de Venta", "CU-23: Emitir comprobante de venta",
     "Comprobar la generación del comprobante para una venta ya pagada, presencial o digital.",
     "La Venta se encuentra en estado \u201cpagada\u201d con su Pago asociado registrado.",
     [
        ("El sistema detecta que la Venta cambió a estado \u201cpagada\u201d.", "Se recopilan los DetalleVenta, el Pago y los datos del cliente."),
        ("Consultar el comprobante desde la pantalla de Caja o desde \u201cMis compras\u201d.", "Se muestra el detalle completo (prendas, montos, pasarela, transacción)."),
        ("Visualizar el comprobante de una venta presencial.", "Incluye los datos del cajero y de la sucursal."),
        ("Visualizar el comprobante de una venta digital pagada con Stripe.", "Incluye el identificador (transaccion_id) del PaymentIntent de Stripe."),
        ("Consultar un comprobante de una venta aún no pagada.", "El sistema no genera el comprobante e informa que la venta no está pagada."),
     ], "Satisfactorio"),

    (30, "Consulta del Historial de Compras", "CU-24: Consultar historial de compras",
     "Validar que el cliente pueda revisar sus compras anteriores, presenciales y digitales.",
     "El cliente ha iniciado sesión y cuenta con al menos una Venta registrada.",
     [
        ("Acceder a \u201cMis compras\u201d.", "El sistema consulta las Venta asociadas al cliente, con sus DetalleVenta."),
        ("Revisar el listado.", "Se muestran tanto compras presenciales como digitales, ordenadas por fechaVenta, indicando tipoOrigen y estado."),
        ("Seleccionar una venta del listado.", "Se muestra el detalle y el comprobante asociado."),
        ("Consultar el historial sin compras registradas.", "El sistema muestra un mensaje informativo en lugar de una lista vacía sin contexto."),
        ("Simular un corte de red al cargar el historial.", "El sistema muestra un mensaje de error tras el tiempo de espera (timeout) en vez de quedar cargando indefinidamente."),
     ], "Satisfactorio"),

    (31, "Interacción con Asistente/Chatbot (IA)", "CU-29: Interactuar con asistente/chatbot (IA)",
     "Verificar que el cliente pueda resolver consultas sobre prendas, tallas, disponibilidad y reservas mediante un asistente conversacional.",
     "El cliente accede a la aplicación web o móvil y el Servicio de IA se encuentra disponible.",
     [
        ("Abrir el asistente y escribir una consulta.", "El sistema complementa la consulta con información de Prenda, VariantePrenda e InventarioSucursal."),
        ("Enviar la consulta al Servicio de IA.", "Se recibe una respuesta en lenguaje natural."),
        ("Verificar el registro de la interacción.", "Queda registrada la InteraccionIA (tipoConsulta=asistente) con el promptConsulta."),
        ("Consultar sobre una prenda concreta.", "La respuesta incluye enlaces al detalle de las prendas mencionadas."),
        ("Realizar una consulta fuera del alcance del sistema.", "El asistente informa la limitación y deriva al catálogo o a la sucursal más cercana."),
     ], "Pendiente \u2014 módulo de IA (recomendaciones/asistente/reportes) aún no implementado en el backend"),

    (32, "Generación de Reporte mediante IA", "CU-30: Generar reporte mediante IA (prompt/voz)",
     "Comprobar que el administrador pueda obtener reportes de negocio formulando la solicitud en lenguaje natural.",
     "El administrador ha iniciado sesión con permisos de reportes y existen Venta, Reserva e InventarioSucursal registrados.",
     [
        ("Seleccionar \u201cGenerar reporte con IA\u201d e ingresar la solicitud por texto.", "El sistema transcribe la solicitud (si fue por voz) y la envía al Servicio de IA."),
        ("El Servicio de IA interpreta la solicitud.", "Devuelve los parámetros del reporte (métricas, periodo y agrupación)."),
        ("El sistema consulta los datos correspondientes.", "Se consultan Venta, DetalleVenta, Reserva e InventarioSucursal según los parámetros obtenidos."),
        ("Mostrar el reporte generado.", "El reporte se presenta al administrador con opción de descarga y se registra la InteraccionIA."),
        ("Enviar una solicitud ambigua al asistente de reportes.", "El sistema informa la ambigüedad y sugiere reformular o usar los reportes predefinidos (CU-31)."),
     ], "Pendiente \u2014 módulo de IA (recomendaciones/asistente/reportes) aún no implementado en el backend"),

    (33, "Visualización del Dashboard de Indicadores", "CU-32: Visualizar dashboard de indicadores",
     "Validar que el administrador pueda visualizar en una sola pantalla los indicadores clave de la operación.",
     "El administrador ha iniciado sesión con permisos de reportes y existe información transaccional registrada.",
     [
        ("Acceder al \u201cDashboard\u201d.", "El sistema calcula los indicadores del periodo vigente a partir de Venta, Reserva e InventarioSucursal."),
        ("Revisar los indicadores mostrados.", "Se muestran las ventas totales, el ticket promedio, las reservas pendientes y las variantes por debajo del stockMinimo."),
        ("Filtrar por Ciudad o Sucursal.", "El sistema recalcula los indicadores según el filtro aplicado."),
        ("Seleccionar un indicador.", "El sistema permite acceder a su reporte detallado (CU-31)."),
        ("Consultar el dashboard en un periodo sin datos.", "El indicador correspondiente se muestra en cero, sin interrumpir la carga del resto del tablero."),
     ], "Pendiente \u2014 módulo de IA/reportes y dashboard aún no implementado en el backend"),

    (34, "Informe de Disponibilidad de Productos por el Proveedor", "CU-34: Informar disponibilidad de productos",
     "Verificar que el proveedor pueda comunicar la disponibilidad y fecha de entrega prevista de sus productos.",
     "El proveedor tiene al menos una Prenda registrada a su nombre (CU-33).",
     [
        ("Acceder a \u201cDisponibilidad de productos\u201d.", "El sistema muestra las Prenda asociadas al Proveedor con su última disponibilidad informada."),
        ("Indicar la cantidad disponible y la fecha estimada de entrega de una prenda.", "El sistema valida que la fecha estimada sea posterior a la fecha actual."),
        ("Guardar la disponibilidad.", "La información queda actualizada y se notifica al administrador."),
        ("Consultar la disponibilidad actualizada desde el panel del administrador.", "El administrador visualiza los datos informados por el proveedor."),
        ("Informar disponibilidad de una prenda dada de baja del catálogo.", "El sistema advierte la situación y no permite registrar la disponibilidad hasta que el administrador reactive la prenda."),
     ], "Pendiente"),
]

PLACEHOLDER_NOTE = (
    "[Espacio reservado para la captura de la evidencia una vez ejecutada la prueba en el ambiente "
    "de pruebas / producción]"
)

def set_cell_shading(cell, color_hex):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color_hex)
    cell._tc.get_or_add_tcPr().append(shd)

def add_header_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        table.rows[i].cells[0].text = k
        table.rows[i].cells[1].text = v
        set_cell_shading(table.rows[i].cells[0], "C6D9F1")
        for p in table.rows[i].cells[0].paragraphs:
            for r in p.runs:
                r.bold = True
    return table

def add_steps_table(doc, steps, estado):
    table = doc.add_table(rows=len(steps) + 1, cols=4)
    table.style = "Table Grid"
    headers = ["Paso", "Acción", "Resultado esperado", "Estado"]
    for j, h in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = h
        set_cell_shading(cell, "C6D9F1")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for i, (accion, resultado) in enumerate(steps, start=1):
        row = table.rows[i]
        row.cells[0].text = str(i)
        row.cells[1].text = accion
        row.cells[2].text = resultado
        row.cells[3].text = estado
    return table

def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = doc.add_heading(
        "FashionStore \u2014 Casos de Prueba Complementarios (CU-04 a CU-34 pendientes)", level=1
    )
    intro = doc.add_paragraph(
        "Este anexo completa la sección 2.5 FT: Pruebas del documento original, agregando los "
        "Casos de Prueba 17 al 34, correspondientes a los casos de uso que no contaban aún con un "
        "caso de prueba documentado (CU-04, CU-05, CU-07, CU-10, CU-15 a CU-24, CU-29, CU-30, CU-32 "
        "y CU-34). Insertar a continuación del \u201cCaso de Prueba 16\u201d y antes del \u201cCódigo QR del Parcial\u201d."
    )
    intro.italic = True

    for num, titulo, caso_uso, descripcion, precond, steps, estado in CASES:
        doc.add_heading(f"Caso de Prueba {num}: {titulo}", level=2)
        add_header_table(doc, [
            ("Caso de uso", caso_uso),
            ("Descripción", descripcion),
            ("Precondiciones", precond),
        ])
        doc.add_paragraph()
        add_steps_table(doc, steps, estado)
        p = doc.add_paragraph()
        run = p.add_run("ADJUNTO PRUEBA IMAGEN")
        run.bold = True
        p2 = doc.add_paragraph(PLACEHOLDER_NOTE)
        p2.italic = True
        doc.add_paragraph()

    out_path = r"c:\Users\AntonioBravo\source\repos\Plataforma-Inteligente-de-comercio-electr-nico\docs\FashionStore_Casos_Prueba_17_34.docx"
    import os
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print("Guardado en:", out_path)

if __name__ == "__main__":
    build()
